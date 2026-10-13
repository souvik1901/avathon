"""
Data Ingestion component (Step 1 of the data-injection pipeline).

Responsibility
--------------
Read every supported file inside a data directory and convert it into the
LangChain ``Document`` data structure. A ``Document`` has two core fields:

* ``page_content`` — the actual text extracted from the file.
* ``metadata``     — extra information about the file (source path, page
  number, file type, ...). Metadata matters because the vector store can
  later *filter* on it during similarity search (e.g. "only search inside
  files authored by X").

Supported formats (the crash course covers PDF/TXT and leaves the rest as an
assignment — all implemented here):

===========  =========================  ==================================
Format       Loader                     One Document per...
===========  =========================  ==================================
``*.pdf``    PyPDFLoader (LangChain)    PDF page
``*.txt``    TextLoader (LangChain)     file
``*.csv``    CSVLoader (LangChain)      row ("column: value" lines)
``*.json``   custom (stdlib ``json``)   FAQ entry ("Q: ...\\nA: ..." text)
``*.xlsx``   custom (``openpyxl``)      sheet row (header mapped to values)
``*.db``     custom (stdlib ``sqlite3``) table row (column: value lines)
``*.docx``   Docx2txtLoader (LangChain) file
===========  =========================  ==================================

The custom loaders show the real lesson of the assignment: *any* data source
can join the pipeline as long as it ends up as Documents. LangChain also
ships ready-made alternatives (``JSONLoader`` — needs the ``jq`` package,
``UnstructuredExcelLoader`` — needs ``unstructured``, ``SQLDatabaseLoader``)
if you prefer not to hand-roll them.
"""

import json
import sqlite3
from pathlib import Path
from typing import List

from langchain_community.document_loaders import (
    CSVLoader,
    Docx2txtLoader,
    PyPDFLoader,
    TextLoader,
)
from langchain_core.documents import Document


# --------------------------------------------------------------------------- #
# Per-format loaders. Each returns a list of Documents whose metadata carries
# `source_file` and `file_type` so every retrieved chunk stays traceable.
# --------------------------------------------------------------------------- #
def _load_pdf_files(data_path: Path) -> List[Document]:
    """Load every ``*.pdf`` with LangChain's PyPDFLoader (one Document per page)."""
    documents: List[Document] = []
    pdf_files = sorted(data_path.glob("**/*.pdf"))
    print(f"Found {len(pdf_files)} PDF files")
    for pdf_file in pdf_files:
        docs = PyPDFLoader(str(pdf_file)).load()
        for doc in docs:
            doc.metadata["source_file"] = pdf_file.name
            doc.metadata["file_type"] = "pdf"
        documents.extend(docs)
        print(f"  loaded {pdf_file.name}: {len(docs)} pages")
    return documents


def _load_text_files(data_path: Path) -> List[Document]:
    """Load every ``*.txt`` with LangChain's TextLoader (one Document per file)."""
    documents: List[Document] = []
    text_files = sorted(data_path.glob("**/*.txt"))
    print(f"Found {len(text_files)} text files")
    for text_file in text_files:
        docs = TextLoader(str(text_file), encoding="utf-8").load()
        for doc in docs:
            doc.metadata["source_file"] = text_file.name
            doc.metadata["file_type"] = "text"
        documents.extend(docs)
        print(f"  loaded {text_file.name}")
    return documents


def _load_csv_files(data_path: Path) -> List[Document]:
    """Load every ``*.csv`` with LangChain's CSVLoader.

    CSVLoader emits **one Document per row**, with the page_content formatted
    as ``column: value`` lines — which reads naturally to both the embedding
    model and the LLM, so no extra processing is needed.
    """
    documents: List[Document] = []
    csv_files = sorted(data_path.glob("**/*.csv"))
    print(f"Found {len(csv_files)} CSV files")
    for csv_file in csv_files:
        docs = CSVLoader(str(csv_file), encoding="utf-8").load()
        for doc in docs:
            doc.metadata["source_file"] = csv_file.name
            doc.metadata["file_type"] = "csv"
        documents.extend(docs)
        print(f"  loaded {csv_file.name}: {len(docs)} rows")
    return documents


def _load_json_files(data_path: Path) -> List[Document]:
    """Load every ``*.json`` into Documents with the stdlib ``json`` module.

    Expects the sample-FAQ shape ``{"faqs": [{question, answer, category}]}``
    and emits one Document per FAQ entry, rendered as ``Q: ...\\nA: ...`` so
    a user question embeds close to the stored question text. Any other JSON
    shape falls back to one Document holding the pretty-printed payload.

    (LangChain's ``JSONLoader`` does the same job generically but requires
    the ``jq`` package for its extraction schema — hand-rolling keeps the
    dependency list short and the mechanics visible.)
    """
    documents: List[Document] = []
    json_files = sorted(data_path.glob("**/*.json"))
    print(f"Found {len(json_files)} JSON files")
    for json_file in json_files:
        payload = json.loads(json_file.read_text(encoding="utf-8"))
        entries = payload.get("faqs") if isinstance(payload, dict) else None
        if isinstance(entries, list):
            for i, entry in enumerate(entries):
                content = f"Q: {entry.get('question', '')}\nA: {entry.get('answer', '')}"
                documents.append(
                    Document(
                        page_content=content,
                        metadata={
                            "source": str(json_file),
                            "source_file": json_file.name,
                            "file_type": "json",
                            "category": entry.get("category", "general"),
                            "entry_index": i,
                        },
                    )
                )
            print(f"  loaded {json_file.name}: {len(entries)} entries")
        else:
            # Unknown shape: keep the whole payload as one Document.
            documents.append(
                Document(
                    page_content=json.dumps(payload, indent=2),
                    metadata={
                        "source": str(json_file),
                        "source_file": json_file.name,
                        "file_type": "json",
                    },
                )
            )
            print(f"  loaded {json_file.name}: 1 document (raw payload)")
    return documents


def _load_excel_files(data_path: Path) -> List[Document]:
    """Load every ``*.xlsx`` into Documents with ``openpyxl``.

    Treats the first row of each sheet as the header and emits **one
    Document per data row**, formatted as ``column: value`` lines (the same
    shape CSVLoader produces, so downstream behaviour is identical).

    (LangChain's ``UnstructuredExcelLoader`` is the ready-made alternative,
    at the cost of the heavyweight ``unstructured`` dependency.)
    """
    from openpyxl import load_workbook

    documents: List[Document] = []
    excel_files = sorted(data_path.glob("**/*.xlsx"))
    print(f"Found {len(excel_files)} Excel files")
    for excel_file in excel_files:
        workbook = load_workbook(str(excel_file), read_only=True, data_only=True)
        file_rows = 0
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                continue
            for i, row in enumerate(rows):
                content = "\n".join(
                    f"{col}: {val}" for col, val in zip(header, row) if val is not None
                )
                if not content:
                    continue
                documents.append(
                    Document(
                        page_content=content,
                        metadata={
                            "source": str(excel_file),
                            "source_file": excel_file.name,
                            "file_type": "excel",
                            "sheet": sheet.title,
                            "row": i + 2,  # 1-based, +1 for the header row
                        },
                    )
                )
                file_rows += 1
        workbook.close()
        print(f"  loaded {excel_file.name}: {file_rows} rows")
    return documents


def _load_sqlite_files(data_path: Path) -> List[Document]:
    """Load every ``*.db`` SQLite database into Documents with ``sqlite3``.

    Iterates every user table (via ``sqlite_master``) and emits **one
    Document per row**, formatted as ``column: value`` lines. The table name
    goes into the metadata so retrieval results can say exactly where a row
    came from.

    (LangChain's ``SQLDatabaseLoader`` covers arbitrary SQL engines through
    SQLAlchemy; stdlib ``sqlite3`` keeps this self-contained.)
    """
    documents: List[Document] = []
    db_files = sorted(data_path.glob("**/*.db"))
    print(f"Found {len(db_files)} SQLite databases")
    for db_file in db_files:
        conn = sqlite3.connect(str(db_file))
        try:
            tables = [
                name
                for (name,) in conn.execute(
                    "SELECT name FROM sqlite_master "
                    "WHERE type='table' AND name NOT LIKE 'sqlite_%'"
                )
            ]
            file_rows = 0
            for table in tables:
                cursor = conn.execute(f'SELECT * FROM "{table}"')
                columns = [desc[0] for desc in cursor.description]
                for i, row in enumerate(cursor):
                    content = "\n".join(
                        f"{col}: {val}" for col, val in zip(columns, row)
                    )
                    documents.append(
                        Document(
                            page_content=content,
                            metadata={
                                "source": str(db_file),
                                "source_file": db_file.name,
                                "file_type": "sqlite",
                                "table": table,
                                "row": i + 1,
                            },
                        )
                    )
                    file_rows += 1
            print(f"  loaded {db_file.name}: {file_rows} rows from {len(tables)} table(s)")
        finally:
            conn.close()
    return documents



def _load_word_files(data_path: Path) -> List[Document]:
    """Load every ``*.docx`` with LangChain's Docx2txtLoader (one Document per file).

    Word documents flatten to plain text, so a whole file arrives as a single
    Document and is chunked downstream like any ``.txt``. Requires the
    ``docx2txt`` package (listed in requirements.txt).
    """
    documents: List[Document] = []
    word_files = sorted(data_path.glob("**/*.docx"))
    print(f"Found {len(word_files)} Word files")
    for word_file in word_files:
        docs = Docx2txtLoader(str(word_file)).load()
        for doc in docs:
            doc.metadata["source_file"] = word_file.name
            doc.metadata["file_type"] = "word"
        documents.extend(docs)
        print(f"  loaded {word_file.name}")
    return documents


# --------------------------------------------------------------------------- #
# Orchestrator
# --------------------------------------------------------------------------- #
def load_all_documents(data_directory: str) -> List[Document]:
    """Load all supported files from ``data_directory`` into Documents.

    Walks the directory recursively and runs every per-format loader
    (PDF, TXT, CSV, JSON, Excel, SQLite, Word). Because each loader outputs the
    same ``Document`` structure, the downstream pipeline — chunking,
    embedding, vector store — needs zero changes when a new format is added.

    Parameters
    ----------
    data_directory:
        Path of the folder that holds the raw knowledge-base files
        (e.g. ``"data"``). Searched recursively.

    Returns
    -------
    list[Document]
        Every page/file/row converted to the LangChain Document structure,
        ready to be passed to the chunking step.
    """
    data_path = Path(data_directory)
    documents: List[Document] = []

    documents.extend(_load_pdf_files(data_path))
    documents.extend(_load_text_files(data_path))
    documents.extend(_load_csv_files(data_path))
    documents.extend(_load_json_files(data_path))
    documents.extend(_load_excel_files(data_path))
    documents.extend(_load_sqlite_files(data_path))
    documents.extend(_load_word_files(data_path))

    print(f"Total documents loaded: {len(documents)}")
    return documents
